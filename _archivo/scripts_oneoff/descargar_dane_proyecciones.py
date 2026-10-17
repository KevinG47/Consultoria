# -*- coding: utf-8 -*-
"""Descarga y describe un cuaderno del DANE de proyecciones departamentales."""
from __future__ import annotations

import io
import ssl
import sys
import urllib.request
from pathlib import Path

import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

RAIZ = Path(__file__).resolve().parents[2]
DESTINO = RAIZ / "_archivo" / "tmp_dane"
DESTINO.mkdir(parents=True, exist_ok=True)

URL = ("https://www.dane.gov.co/files/censo2018/proyecciones-de-poblacion/"
       "Departamental/PPED-AreaDep-2018-2050_VP.xlsx")
ARCHIVO = DESTINO / URL.rsplit("/", 1)[-1]
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) consultoria-usta/1.0"}

if not ARCHIVO.exists():
    print(f"Descargando {URL}")
    req = urllib.request.Request(URL, headers=UA)
    with urllib.request.urlopen(req, timeout=180,
                                context=ssl.create_default_context()) as r:
        ARCHIVO.write_bytes(r.read())
print(f"{ARCHIVO.name}: {ARCHIVO.stat().st_size:,} bytes")

wb = openpyxl.load_workbook(ARCHIVO, read_only=True, data_only=True)
print(f"Hojas: {wb.sheetnames}")
for nombre in wb.sheetnames:
    ws = wb[nombre]
    print(f"\n=== {nombre!r}  filas={ws.max_row} col={ws.max_column} ===")
    for i, fila in enumerate(ws.iter_rows(min_row=1, max_row=18,
                                          max_col=14, values_only=True), start=1):
        celdas = ["" if c is None else str(c)[:16] for c in fila]
        print(f"  {i:3d} | " + " | ".join(celdas))
wb.close()
