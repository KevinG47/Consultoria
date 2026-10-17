# -*- coding: utf-8 -*-
"""Vuelca el índice del cuaderno DANE y la estructura de la hoja departamental 1PD."""
from __future__ import annotations

import io
import sys
from pathlib import Path

import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

RAIZ = Path(__file__).resolve().parents[2]
XLSX = (RAIZ / "_archivo" / "tmp_dane"
        / "PERSONAS_DEMOGRAFICO_Cuadros_CNPV_2018.xlsx")

wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)

print("=" * 100)
print("ÍNDICE COMPLETO")
print("=" * 100)
ws = wb["Índice"]
for fila in ws.iter_rows(values_only=True):
    celdas = [str(c) for c in fila if c is not None and str(c).strip()]
    if celdas:
        print("  " + "  |  ".join(c[:120] for c in celdas))

print()
print("=" * 100)
print("HOJA 1PD — cabecera y filas 'Total' por departamento")
print("=" * 100)
ws = wb["1PD"]
filas = list(ws.iter_rows(min_row=1, max_row=60, values_only=True))
for i, fila in enumerate(filas, start=1):
    celdas = [("" if c is None else str(c)[:30]) for c in fila[:6]]
    print(f"  {i:4d} | " + " | ".join(celdas))

print("\n--- todas las filas con etiqueta 'Total' en alguna de las 3 primeras columnas ---")
ws = wb["1PD"]
n = 0
for i, fila in enumerate(ws.iter_rows(values_only=True), start=1):
    etiquetas = [str(c).strip() if c is not None else "" for c in fila[:3]]
    if any(e.lower() == "total" for e in etiquetas):
        n += 1
        if n <= 45:
            print(f"  {i:5d} | " + " | ".join(f"{e[:34]:34s}" for e in etiquetas)
                  + f" || tot={fila[3]}")
print(f"  -> {n} filas con 'Total'")

wb.close()
